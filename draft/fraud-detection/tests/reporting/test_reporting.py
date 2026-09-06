"""Unit tests for the modularized reporting pipeline."""

from __future__ import annotations

from docx import Document

from src.reporting.docx_table_builder import add_spacer, add_table
from src.reporting.markdown_parser import add_runs, clean_markdown_symbols, parse_markdown
from src.reporting.toc_generator import collect_toc_entries


def test_clean_markdown_symbols() -> None:
    assert clean_markdown_symbols("`code` with $math$") == "code with math"
    assert clean_markdown_symbols("`amount_to_oldbalance_ratio`") == "amount_to_oldbalance_ratio"


def test_parse_markdown_blocks() -> None:
    md_text = """
# Tiêu đề 1

Đoạn văn thường với **in đậm**.

| Cột 1 | Cột 2 |
|---|---|
| A | B |

- Mục 1
- Mục 2

---
"""
    blocks = parse_markdown(md_text)
    types = [b["type"] for b in blocks]
    assert "h1" in types
    assert "para" in types
    assert "table" in types
    assert "bullet" in types
    assert "pagebreak" in types


def test_collect_toc_entries() -> None:
    blocks = [
        {"type": "h1", "text": "Bìa Báo Cáo"},
        {"type": "h1", "text": "CHƯƠNG 1. GIỚI THIỆU"},
        {"type": "h2", "text": "1.1. Bối cảnh"},
        {"type": "h2", "text": "Mục không đánh số"},
        {"type": "h3", "text": "1.1.1. Chi tiết"},
    ]
    entries = collect_toc_entries(blocks)
    assert len(entries) == 2
    assert entries[0] == {"text": "CHƯƠNG 1. GIỚI THIỆU", "level": 1}
    assert entries[1] == {"text": "1.1. Bối cảnh", "level": 2}


def test_add_table_and_runs() -> None:
    doc = Document()
    p = doc.add_paragraph()
    add_runs(p, "Văn bản **đậm** thường")
    assert len(p.runs) == 3
    assert p.runs[1].bold is True

    table = add_table(doc, ["STT", "Tên"], [["1", "Thử nghiệm"]])
    assert len(table.rows) == 2
    assert len(table.columns) == 2
    assert table.style.name == "Table Grid"

    spacer = add_spacer(doc)
    assert spacer.paragraph_format.line_spacing == 1.0
