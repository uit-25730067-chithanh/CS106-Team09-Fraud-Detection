"""Phase 07: Dựng bản Word của báo cáo từ `reports/report-source.md`.

Script giữ đúng định dạng đã chốt ở Sprint 2 bằng cách dùng
`reports/_report_template.docx` làm khuôn: template chỉ chứa styles, header,
footer, khổ giấy A4 và lề, phần thân hoàn toàn rỗng. Nhờ vậy mỗi lần source
thay đổi chỉ cần chạy lại script, không phải chỉnh tay trong Word.

Mục lục được sinh tự động từ các heading cấp 1 và cấp 2. Số trang lấy từ
`reports/_toc_pages.json`. Tệp đó do `sync_toc_pages.py` đọc ngược từ bản PDF
đã xuất, nên quy trình đầy đủ là: dựng DOCX, xuất PDF, đồng bộ số trang, dựng
lại DOCX rồi xuất PDF lần cuối.

Cách dùng:
    python build_report.py

Author: Vũ Văn Duy
"""
import json
import os
import re
import sys

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

REPORTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "reports")
SOURCE_MD = os.path.join(REPORTS_DIR, "report-source.md")
TEMPLATE = os.path.join(REPORTS_DIR, "_report_template.docx")
TOC_PAGES = os.path.join(REPORTS_DIR, "_toc_pages.json")
OUTPUT = os.path.join(REPORTS_DIR, "[Nhom9]_BaoCao_FraudDetection_Sprint3_Duy.docx")

# Bề rộng vùng in: khổ A4 11906 twips trừ lề trái 1701 và lề phải 1134.
CONTENT_WIDTH_TWIPS = 9071
NARROW_HEADERS = {"#", "STT", "Hạng"}
NARROW_WIDTH_TWIPS = 700
TABLE_FONT_PT = 9.5
TOC_FIRST_PAGE_ROWS = 20
TOC_PAGE_COLUMN_TWIPS = 648
FIGURE_WIDTH_TWIPS = 8600
REFERENCE_HANGING_TWIPS = 567


# ─────────────────────────── Phân tích Markdown ────────────────────────────

def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown(text: str) -> list[dict]:
    """Chuyển Markdown của báo cáo thành danh sách block đơn giản."""

    lines = text.splitlines()
    blocks: list[dict] = []
    index = 0

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            blocks.append({"type": "pagebreak"})
            index += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            blocks.append({
                "type": f"h{len(heading.group(1))}",
                "text": heading.group(2).strip(),
            })
            index += 1
            continue

        image = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if image:
            blocks.append({
                "type": "image",
                "caption": image.group(1).strip(),
                "path": image.group(2).strip(),
            })
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\|[\s:|-]+\|$", lines[index + 1].strip()
        ):
            header = _split_row(stripped)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_row(lines[index].strip()))
                index += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        if stripped.startswith("> "):
            note = [stripped[2:].strip()]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("> "):
                note.append(lines[index].strip()[2:].strip())
                index += 1
            blocks.append({"type": "note", "text": " ".join(note)})
            continue

        if stripped.startswith("- "):
            blocks.append({"type": "bullet", "text": stripped[2:].strip()})
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            blocks.append({"type": "number", "text": numbered.group(1).strip()})
            index += 1
            continue

        if re.match(r"^\*\*.*\*\*$", stripped):
            blocks.append({"type": "caption", "text": stripped.strip("*").strip()})
            index += 1
            continue

        blocks.append({"type": "para", "text": stripped})
        index += 1

    return blocks


# ─────────────────────────── Định dạng nội tuyến ───────────────────────────

def _clean(text: str) -> str:
    """Bỏ ký hiệu Markdown mà bản Word không hiển thị (backtick, ký hiệu toán)."""

    text = text.replace("`", "")
    return re.sub(r"\$(.+?)\$", r"\1", text)


def add_runs(paragraph, text: str, bold: bool = False, size_pt: float | None = None):
    """Ghi text vào paragraph, tách phần **in đậm** thành run riêng."""

    for segment in re.split(r"(\*\*.+?\*\*)", _clean(text)):
        if not segment:
            continue
        run = paragraph.add_run(segment.strip("*") if segment.startswith("**") else segment)
        run.bold = bold or segment.startswith("**")
        if size_pt is not None:
            run.font.size = Pt(size_pt)
    return paragraph


# ───────────────────────────── Dựng bảng Word ──────────────────────────────

def _column_widths(header: list[str], rows: list[list[str]]) -> list[int]:
    """Chia bề rộng theo độ dài nội dung, cột số thứ tự giữ hẹp cố định."""

    weights = []
    for column in range(len(header)):
        cells = [header[column]] + [row[column] for row in rows if column < len(row)]
        weights.append(max(len(_clean(cell)) for cell in cells))

    narrow = header[0] in NARROW_HEADERS
    available = CONTENT_WIDTH_TWIPS - (NARROW_WIDTH_TWIPS if narrow else 0)
    scalable = weights[1:] if narrow else weights
    total = sum(scalable) or 1
    widths = [round(available * weight / total) for weight in scalable]
    return ([NARROW_WIDTH_TWIPS] + widths) if narrow else widths


def _is_numeric_column(rows: list[list[str]], column: int) -> bool:
    cells = [row[column] for row in rows if column < len(row)]
    if not cells:
        return False
    numeric = sum(bool(re.fullmatch(r"[\d.,%\s]+", _clean(cell).strip("*"))) for cell in cells)
    return numeric >= max(1, round(0.6 * len(cells)))


def add_spacer(document):
    """Đệm mỏng sau bảng: đủ tách khỏi đoạn kế tiếp, không tạo dòng trống thừa."""

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 1.0
    spacer.add_run().font.size = Pt(6)
    return spacer


def _repeat_header(table) -> None:
    """Lặp lại hàng tiêu đề khi bảng bị tách sang trang sau."""

    # Thứ tự tblHeader trước cantSplit lấy theo đúng bản Sprint 2 — Pages chỉ
    # lặp lại hàng tiêu đề khi trPr được viết theo thứ tự này.
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_pr.append(header_pr.makeelement(qn("w:tblHeader"), {qn("w:val"): "true"}))
    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        row_pr.append(row_pr.makeelement(qn("w:cantSplit"), {}))


def _fix_layout(table, widths: list[int]) -> None:
    """Khoá bề rộng cột.

    Chỉ gán ``cell.width`` là chưa đủ: trình soạn thảo vẫn co giãn theo
    ``tblGrid``. Phải đặt ``tblLayout`` cố định và ghi bề rộng vào cả cột.
    """

    # ``autofit = False`` đã ghi <w:tblLayout w:type="fixed"/> đúng vị trí trong
    # tblPr. Tự append thêm một thẻ nữa sẽ sai thứ tự lược đồ và làm Pages bỏ
    # qua các thuộc tính phía sau, trong đó có lặp lại hàng tiêu đề.
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Twips(width)


def add_table(document, header: list[str], rows: list[list[str]]):
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "TableGrid"

    widths = _column_widths(header, rows)
    _fix_layout(table, widths)
    _repeat_header(table)
    narrow_first = header[0] in NARROW_HEADERS
    alignments = [
        WD_ALIGN_PARAGRAPH.CENTER if narrow_first else WD_ALIGN_PARAGRAPH.LEFT
    ] + [
        WD_ALIGN_PARAGRAPH.CENTER if _is_numeric_column(rows, column) else WD_ALIGN_PARAGRAPH.LEFT
        for column in range(1, len(header))
    ]

    for column, title in enumerate(header):
        cell = table.cell(0, column)
        cell.width = Twips(widths[column])
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignments[column]
        add_runs(paragraph, title, bold=True, size_pt=TABLE_FONT_PT)

    for row_index, row in enumerate(rows, start=1):
        for column in range(len(header)):
            cell = table.cell(row_index, column)
            cell.width = Twips(widths[column])
            paragraph = cell.paragraphs[0]
            paragraph.alignment = alignments[column]
            add_runs(paragraph, row[column] if column < len(row) else "",
                     size_pt=TABLE_FONT_PT)

    return table


# ────────────────────────────── Mục lục ────────────────────────────────────

def collect_toc_entries(blocks: list[dict]) -> list[dict]:
    """Mục lục gồm tiêu đề chương và các mục đánh số dạng x.y.

    Bỏ qua tiêu đề bìa (block đầu tiên), heading cấp 3 và các mục không đánh số
    trong phần tài liệu tham khảo.
    """

    entries = []
    for index, block in enumerate(blocks):
        if index == 0 or block["type"] not in ("h1", "h2"):
            continue
        if block["type"] == "h2" and not re.match(r"^\d+\.\d+\.", block["text"]):
            continue
        entries.append({"text": block["text"], "level": 1 if block["type"] == "h1" else 2})
    return entries


def add_toc(document, entries: list[dict], pages: dict[str, int]) -> None:
    chunks = [entries[:TOC_FIRST_PAGE_ROWS], entries[TOC_FIRST_PAGE_ROWS:]]
    for chunk_index, chunk in enumerate(chunks):
        if not chunk:
            continue
        if chunk_index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph("MỤC LỤC (TIẾP THEO)" if chunk_index else "MỤC LỤC",
                               style="Heading1")

        table = document.add_table(rows=len(chunk), cols=2)
        _fix_layout(table, [CONTENT_WIDTH_TWIPS - TOC_PAGE_COLUMN_TWIPS,
                            TOC_PAGE_COLUMN_TWIPS])
        for row_index, entry in enumerate(chunk):
            label_cell, page_cell = table.cell(row_index, 0), table.cell(row_index, 1)
            label_cell.width = Twips(CONTENT_WIDTH_TWIPS - TOC_PAGE_COLUMN_TWIPS)
            page_cell.width = Twips(TOC_PAGE_COLUMN_TWIPS)

            label = label_cell.paragraphs[0]
            label.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if entry["level"] == 2:
                label.paragraph_format.left_indent = Twips(280)
            run = label.add_run(_clean(entry["text"]))
            run.bold = entry["level"] == 1
            run.font.size = Pt(11)

            page = page_cell.paragraphs[0]
            page.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            number = pages.get(entry["text"])
            page_run = page.add_run(str(number) if number else "")
            page_run.bold = entry["level"] == 1
            page_run.font.size = Pt(11)


# ─────────────────────────────── Dựng file ─────────────────────────────────

def build(blocks: list[dict], pages: dict[str, int]) -> Document:
    document = Document(TEMPLATE)
    entries = collect_toc_entries(blocks)
    cover_done = False

    for index, block in enumerate(blocks):
        kind = block["type"]

        if kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if not cover_done:
                add_toc(document, entries, pages)
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                cover_done = True
            continue

        if kind == "h1":
            if index == 0:
                paragraph = document.add_paragraph(_clean(block["text"]), style="CoverTitle")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                document.add_paragraph(_clean(block["text"]), style="Heading1")
            continue

        if kind in ("h2", "h3"):
            document.add_paragraph(_clean(block["text"]),
                                   style="Heading2" if kind == "h2" else "Heading3")
            continue

        if kind == "note":
            paragraph = document.add_paragraph(style="ReportNote")
            add_runs(paragraph, block["text"])
            continue

        if kind == "caption":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            add_runs(paragraph, block["text"], bold=True)
            continue

        if kind == "table":
            add_table(document, block["header"], block["rows"])
            add_spacer(document)
            continue

        if kind == "image":
            path = os.path.join(REPORTS_DIR, block["path"])
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.add_run().add_picture(path, width=Twips(FIGURE_WIDTH_TWIPS))
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.add_run(_clean(block["caption"])).italic = True
            continue

        if kind in ("bullet", "number"):
            paragraph = document.add_paragraph(
                style="ListBullet" if kind == "bullet" else "ListNumber")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(paragraph, block["text"])
            continue

        style = "CoverMeta" if not cover_done else "ReportBody"
        paragraph = document.add_paragraph(style=style)
        if not cover_done:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r"^\[\d+\]\s", block["text"]):
            paragraph.paragraph_format.left_indent = Twips(REFERENCE_HANGING_TWIPS)
            paragraph.paragraph_format.first_line_indent = Twips(-REFERENCE_HANGING_TWIPS)
            # Căn đều làm giãn ký tự ở dòng chứa DOI dài không ngắt được.
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_runs(paragraph, block["text"])

    return document


def main() -> None:
    print("=== Phase 07: Dựng bản Word từ report-source.md ===")

    with open(SOURCE_MD, encoding="utf-8") as f:
        blocks = parse_markdown(f.read())

    pages: dict[str, int] = {}
    if os.path.exists(TOC_PAGES):
        with open(TOC_PAGES, encoding="utf-8") as f:
            pages = json.load(f)
        print(f"Đã nạp số trang mục lục: {len(pages)} mục")
    else:
        print("Chưa có _toc_pages.json — mục lục sẽ để trống số trang")

    document = build(blocks, pages)
    document.save(OUTPUT)

    counts = {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "toc entries": len(collect_toc_entries(blocks)),
    }
    print("Saved -> reports/" + os.path.basename(OUTPUT))
    print("  " + " | ".join(f"{key}: {value}" for key, value in counts.items()))


if __name__ == "__main__":
    main()
