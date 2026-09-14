"""Document Body Builder (Section 3) for UIT Academic Documents.

Renders the main body of the report from Markdown source lines:
- Margins: Left 3.0cm, Right 2.0cm, Top 2.0cm, Bottom 2.0cm
- Font: Times New Roman 13pt, 1.3 line spacing, 1.27cm first-line indent
- Justified alignment, 100% pure black text (#000000)
- Bookmark anchors on all headings (matching TOC)
- Active in-text citations jumping to bibliography
- Tables rendered via ``add_styled_table``
"""

from __future__ import annotations

import os
import re

import docx
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt

from .docx_table_builder import add_styled_table
from .hyperlink_helpers import (
    COLOR_TEXT_MAIN,
    add_bookmark,
    add_page_number_field,
)
from .inline_parser import parse_inline_formatting


def build_body_section(
    doc: docx.Document,
    lines: list[str],
    meta: dict,
    base_dir: str,
    toc_data: list[dict],
) -> None:
    """Build Section 3 (Document Body) adhering to UIT standards."""
    sec3 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec3.page_width = Cm(21.0)
    sec3.page_height = Cm(29.7)
    sec3.top_margin = Cm(2.0)
    sec3.bottom_margin = Cm(2.0)
    sec3.left_margin = Cm(3.0)
    sec3.right_margin = Cm(2.0)
    sec3.header_distance = Cm(1.0)
    sec3.footer_distance = Cm(1.0)
    sec3.different_first_page_header_footer = False
    sec3.header.is_linked_to_previous = False
    sec3.footer.is_linked_to_previous = False

    # Remove borders
    sectPr3 = sec3._sectPr
    pgBorders3 = parse_xml(
        r"""
    <w:pgBorders %s w:offsetFrom="page">
        <w:top w:val="none"/><w:left w:val="none"/>
        <w:bottom w:val="none"/><w:right w:val="none"/>
    </w:pgBorders>
    """
        % nsdecls("w")
    )
    sectPr3.append(pgBorders3)

    # Page number starts at 1
    pgNumType = parse_xml(
        r'<w:pgNumType %s w:start="1"/>' % nsdecls("w")
    )
    sectPr3.append(pgNumType)

    # Header running title
    hdr_p = sec3.header.paragraphs[0]
    hdr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr_p.paragraph_format.space_after = Pt(0)
    hdr_r = hdr_p.add_run(
        meta.get("header_title", "Báo cáo Đồ án CS106 | Nhóm 9 - Fraud Detection")
    )
    hdr_r.font.name = "Times New Roman"
    hdr_r.font.size = Pt(9)
    hdr_r.italic = True
    hdr_r.font.color.rgb = COLOR_TEXT_MAIN

    # Footer page number
    ftr_p = sec3.footer.paragraphs[0]
    ftr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ftr_p.paragraph_format.space_after = Pt(0)
    ftr_p.paragraph_format.space_before = Pt(0)
    ftr_r = ftr_p.add_run()
    ftr_r.font.name = "Times New Roman"
    ftr_r.font.size = Pt(11)
    ftr_r.font.color.rgb = COLOR_TEXT_MAIN
    add_page_number_field(ftr_r)

    # Build title-to-anchor registry
    anchor_by_title = {
        item["text"].strip().lower(): item.get("anchor") for item in toc_data
    }

    # Skip draft metadata header if source starts with # Title / metadata table
    start_idx = 0
    for idx, line in enumerate(lines):
        if line.strip().startswith("# TÓM TẮT") or line.strip().startswith(
            "# CHƯƠNG 1"
        ):
            start_idx = idx
            break

    i = start_idx
    bookmark_counter = 1000

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 1. Skip dividers, page breaks, empty lines
        if (
            stripped in ["---", "***", "___", "\\newpage", "<br>", "<br><br>"]
            or not stripped
        ):
            i += 1
            continue

        # 2. Image tag: ![alt](path)
        img_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if img_match:
            bookmark_counter, i = _handle_image(
                doc, img_match, base_dir, bookmark_counter, i
            )
            continue

        # 3. Figure caption standalone: *Hình X: ...*
        if re.match(r"^\*Hình\s+\d+.*\*$", stripped):
            _add_figure_caption(doc, stripped[1:-1].strip())
            i += 1
            continue

        # 4–6. Headings
        if line.startswith("# "):
            bookmark_counter = _handle_heading(
                doc, line[2:], 1, anchor_by_title, bookmark_counter
            )
            i += 1
            continue
        if line.startswith("## "):
            bookmark_counter = _handle_heading(
                doc, line[3:], 2, anchor_by_title, bookmark_counter
            )
            i += 1
            continue
        if line.startswith("### "):
            bookmark_counter = _handle_heading(
                doc, line[4:], 3, anchor_by_title, bookmark_counter
            )
            i += 1
            continue

        # 7. Blockquote: > ...
        if line.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q_text = lines[i].strip()[1:].strip()
                if q_text:
                    quote_lines.append(q_text)
                i += 1
            full_quote = " ".join(quote_lines)
            p_q = doc.add_paragraph()
            p_q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_q.paragraph_format.left_indent = Cm(1.27)
            p_q.paragraph_format.right_indent = Cm(0.8)
            p_q.paragraph_format.space_before = Pt(3)
            p_q.paragraph_format.space_after = Pt(4)
            p_q.paragraph_format.line_spacing = 1.2
            parse_inline_formatting(
                p_q, full_quote, is_italic=True, base_font_size=12.0
            )
            continue

        # 8. List items (*, -, +, or 1.)
        list_match = re.match(r"^(\s*)([\*\-\+]|\d+\.)\s+(.*)$", line)
        if list_match:
            _handle_list_item(doc, list_match)
            i += 1
            continue

        # 9. Reference / Bibliography entry: [1] Author...
        ref_match = re.match(r"^\[(\d+)\]\s+(.*)$", stripped)
        if ref_match:
            bookmark_counter = _handle_reference(
                doc, ref_match, bookmark_counter
            )
            i += 1
            continue

        # 10. Markdown Table: | col1 | col2 |
        if stripped.startswith("|") and "|" in stripped[1:]:
            i = _handle_table(doc, lines, i)
            continue

        # 11. Standalone Table Caption: **Bảng X.Y...** or *Bảng X.Y...*
        if re.match(r"^(\*\*|\*)?Bảng\s+\d+.*?(\*\*|\*)?$", stripped):
            _add_table_caption(doc, stripped)
            i += 1
            continue

        # 12. Standard Body Paragraph
        p_body = doc.add_paragraph()
        p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_body.paragraph_format.first_line_indent = Cm(1.27)
        p_body.paragraph_format.space_before = Pt(0)
        p_body.paragraph_format.space_after = Pt(2.0)
        p_body.paragraph_format.line_spacing = 1.3
        parse_inline_formatting(p_body, stripped)
        i += 1


# ---------------------------------------------------------------------------
# Private helpers — keep the main loop readable
# ---------------------------------------------------------------------------


def _handle_image(doc, img_match, base_dir, bookmark_counter, i):
    """Insert an image and optional caption."""
    img_caption = img_match.group(1).strip()
    img_rel_path = img_match.group(2).strip()
    img_full = os.path.normpath(os.path.join(base_dir, img_rel_path))

    if os.path.exists(img_full):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        r_img = p_img.add_run()
        r_img.add_picture(img_full, width=Cm(14.5))

        if img_caption:
            _add_figure_caption(doc, img_caption)
    return bookmark_counter, i + 1


def _add_figure_caption(doc, caption_text: str) -> None:
    """Add a centered italic figure caption."""
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10.5)
    r_cap.italic = True
    r_cap.font.color.rgb = COLOR_TEXT_MAIN


_HEADING_CONFIG = {
    1: {"font_size": 15.0, "space_before": 24, "space_after": 12, "italic": False},
    2: {"font_size": 14.0, "space_before": 18, "space_after": 6, "italic": False},
    3: {"font_size": 13.0, "space_before": 12, "space_after": 6, "italic": True},
}


def _handle_heading(
    doc, raw_text: str, level: int, anchor_by_title: dict, bookmark_counter: int
) -> int:
    """Add a heading paragraph with optional bookmark anchor."""
    h_text = raw_text.strip()
    cfg = _HEADING_CONFIG[level]

    p_h = doc.add_paragraph()
    p_h.paragraph_format.keep_with_next = True
    p_h.paragraph_format.space_before = Pt(cfg["space_before"])
    p_h.paragraph_format.space_after = Pt(cfg["space_after"])
    p_h.paragraph_format.line_spacing = 1.15
    p_h.paragraph_format.first_line_indent = Cm(0)
    if level == 1:
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p_h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    anchor = anchor_by_title.get(h_text.lower())
    if anchor:
        bookmark_counter += 1
        add_bookmark(p_h, anchor, bookmark_counter)

    r = p_h.add_run(h_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(cfg["font_size"])
    r.bold = True
    r.italic = cfg["italic"]
    r.font.color.rgb = COLOR_TEXT_MAIN

    return bookmark_counter


def _handle_list_item(doc, list_match) -> None:
    """Add a bullet or numbered list item.

    Danh sách căn trái thay vì căn đều. Các mục thường chứa đường dẫn tệp hoặc
    tên định danh dài không có chỗ ngắt dòng, nên căn đều sẽ kéo giãn ký tự của
    những dòng còn lại trong cùng mục.
    """
    indent_spaces = len(list_match.group(1))
    bullet_sym = list_match.group(2)
    content = list_match.group(3)

    p_l = doc.add_paragraph()
    p_l.alignment = WD_ALIGN_PARAGRAPH.LEFT
    level = indent_spaces // 2
    p_l.paragraph_format.left_indent = Cm(1.27 + level * 0.4)
    p_l.paragraph_format.first_line_indent = Cm(-0.5)
    p_l.paragraph_format.space_before = Pt(1)
    p_l.paragraph_format.space_after = Pt(2)
    p_l.paragraph_format.line_spacing = 1.25

    bullet_char = "-  " if bullet_sym in ["*", "-", "+"] else f"{bullet_sym} "
    r_b = p_l.add_run(bullet_char)
    r_b.font.name = "Times New Roman"
    r_b.font.size = Pt(13)
    r_b.bold = True
    r_b.font.color.rgb = COLOR_TEXT_MAIN

    parse_inline_formatting(p_l, content)


def _handle_reference(doc, ref_match, bookmark_counter: int) -> int:
    """Add a bibliography reference entry with bookmark anchor."""
    ref_num = ref_match.group(1)
    ref_body = ref_match.group(2)

    p_ref = doc.add_paragraph()
    p_ref.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ref.paragraph_format.left_indent = Cm(0.8)
    p_ref.paragraph_format.first_line_indent = Cm(-0.8)
    p_ref.paragraph_format.space_before = Pt(2)
    p_ref.paragraph_format.space_after = Pt(3)
    p_ref.paragraph_format.line_spacing = 1.15

    bookmark_counter += 1
    add_bookmark(p_ref, f"ref_{ref_num}", bookmark_counter)

    r_num = p_ref.add_run(f"[{ref_num}]\t")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(11.5)
    r_num.bold = True
    r_num.font.color.rgb = COLOR_TEXT_MAIN

    parse_inline_formatting(p_ref, ref_body, base_font_size=11.5)
    return bookmark_counter


def _handle_table(doc, lines: list[str], i: int) -> int:
    """Parse a markdown table and render it via add_styled_table."""
    table_lines = []
    while (
        i < len(lines)
        and lines[i].strip().startswith("|")
        and "|" in lines[i].strip()[1:]
    ):
        table_lines.append(lines[i].strip())
        i += 1

    rows_data = []
    for t_line in table_lines:
        if re.match(r"^\|[\s\:\-]+(\Auto|[\s\:\-\|]+)\|$", t_line):
            continue
        raw_cells = [c.strip() for c in t_line.strip("|").split("|")]
        rows_data.append(raw_cells)

    if rows_data:
        header = rows_data[0]
        body_rows = rows_data[1:]

        cap = None
        # Case A: Caption placed before table
        if len(doc.paragraphs) > 0:
            last_p_text = doc.paragraphs[-1].text.strip()
            if last_p_text.startswith("Bảng ") or last_p_text.startswith("Table "):
                cap = last_p_text
                p_elem = doc.paragraphs[-1]._p
                p_elem.getparent().remove(p_elem)

        # Case B: Caption placed directly after table
        if not cap and i < len(lines):
            next_stripped = lines[i].strip()
            if re.match(
                r"^(\*\*|\*)?(Bảng|Table)\s+\d+.*(\*\*|\*)?$", next_stripped
            ):
                cap = re.sub(r"[*_]", "", next_stripped).strip()
                i += 1

        add_styled_table(
            doc,
            header,
            body_rows,
            caption=cap,
            parse_inline_func=lambda p, txt, **kw: parse_inline_formatting(
                p, txt, **kw
            ),
        )
    return i


def _add_table_caption(doc, stripped: str) -> None:
    """Add a standalone table caption (centered, italic)."""
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    p_cap.paragraph_format.line_spacing = 1.15
    clean_cap = re.sub(r"[*_]", "", stripped).strip()
    r_cap = p_cap.add_run(clean_cap)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10.5)
    r_cap.italic = True
    r_cap.bold = False
    r_cap.font.color.rgb = COLOR_TEXT_MAIN
