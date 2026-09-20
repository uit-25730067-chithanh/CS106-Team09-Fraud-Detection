"""
Script tạo file Excel Danh sách nhóm theo chuẩn yêu cầu của Giảng viên môn CS106 - UIT.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_team_excel(output_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách nhóm 09"
    ws.views.sheetView[0].showGridLines = True

    # Styles
    font_title = Font(name="Calibri", size=14, bold=True, color="1F497D")
    font_sub = Font(name="Calibri", size=11, bold=True, color="333333")
    font_topic = Font(name="Calibri", size=11, italic=True, color="555555")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=11, color="000000")
    font_leader = Font(name="Calibri", size=11, bold=True, color="000000")

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color="D9D9D9"),
        right=Side(style='thin', color="D9D9D9"),
        top=Side(style='thin', color="D9D9D9"),
        bottom=Side(style='thin', color="D9D9D9")
    )
    header_border = Border(
        left=Side(style='thin', color="1F497D"),
        right=Side(style='thin', color="1F497D"),
        top=Side(style='medium', color="1F497D"),
        bottom=Side(style='medium', color="1F497D")
    )

    # Title lines
    ws.merge_cells("A1:F1")
    ws["A1"] = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN — ĐHQG-HCM"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:F2")
    ws["A2"] = "MÔN HỌC: TRÍ TUỆ NHÂN TẠO (CS106) — HỌC KỲ 2, NĂM HỌC 2025-2026"
    ws["A2"].font = font_sub
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:F3")
    ws["A3"] = "ĐỀ TÀI #7: PHÁT HIỆN GIAO DỊCH TÀI CHÍNH BẤT THƯỜNG VÀ NGHI VẤN GIAN LẬN (FRAUD DETECTION)"
    ws["A3"].font = font_topic
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A4:F4")
    ws["A4"] = "BẢNG TỔNG HỢP DANH SÁCH & PHÂN CÔNG NHIỆM VỤ THÀNH VIÊN — NHÓM 09"
    ws["A4"].font = Font(name="Calibri", size=12, bold=True, color="C00000")
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 10  # blank line

    # Table headers
    headers = [
        ("STT", Alignment(horizontal="center", vertical="center")),
        ("Họ và tên", Alignment(horizontal="left", vertical="center")),
        ("MSSV", Alignment(horizontal="center", vertical="center")),
        ("Lớp", Alignment(horizontal="center", vertical="center")),
        ("Vai trò chính", Alignment(horizontal="left", vertical="center")),
        ("Phân công chi tiết (Tasks đảm nhiệm)", Alignment(horizontal="left", vertical="center"))
    ]

    start_row = 6
    ws.row_dimensions[start_row].height = 28

    for col_idx, (header_text, align) in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = header_border

    # Data rows
    members = [
        (1, "Trần Hoàng Hôn", "26410046", "A", "Nhóm trưởng / Project Manager", "Quản lý tiến độ tổng thể, Chủ trì slide thuyết trình PPT, Điều phối kịch bản bảo vệ, Đóng gói nộp bài"),
        (2, "Nguyễn Duy Khang", "26410055", "A", "Thành viên / Evaluation Specialist", "Hiện thực scripts đánh giá (Confusion Matrix, ROC/PR Curves), So sánh chéo 5 mô hình (Phase 05), Kiểm định thống kê"),
        (3, "Vũ Văn Duy", "26410031", "A", "Thành viên / Lead Technical Writer", "Chủ trì soạn thảo Báo cáo đồ án khoa học (Word/PDF 32 trang), Tổng hợp cơ sở lý thuyết, Viết phân tích thảo luận"),
        (4, "Phạm Thành Trung", "26410141", "A", "Thành viên / UI/UX & Demo Engineer", "Xây dựng ứng dụng Demo Web (Streamlit UI), Tích hợp mô hình suy luận XGBoost, Chuyển đổi định dạng VNĐ & PaySim, Browser QA"),
        (5, "Đặng Chí Thanh", "25730067", "B", "Thành viên / Data & Pipeline Engineer", "Phân tích thăm dò dữ liệu (EDA), Xây dựng Data Loader, Feature Scaler (14 đặc trưng), Derived features, Kiểm định không rò rỉ dữ liệu"),
        (6, "Bùi Thị Mỷ Cẩm", "25730013", "B", "Thành viên / ML & Deep Learning Engineer", "Huấn luyện & tối ưu mô hình XGBoost (SMOTE/ADASYN), Xây dựng Autoencoder phát hiện bất thường không giám sát, Xuất Artifacts"),
        (7, "Hoàng Cao Sơn", "25730061", "B", "Thành viên / Imbalance & ML Specialist", "Phân tích mất cân bằng dữ liệu, Áp dụng SMOTENC và ADASYN bảo toàn biến nhị phân, Huấn luyện mô hình Random Forest, Phân tích độ quan trọng đặc trưng")
    ]

    for row_idx, member in enumerate(members, start=start_row + 1):
        ws.row_dimensions[row_idx].height = 36
        is_leader = (member[0] == 1)
        use_zebra = (row_idx % 2 == 1)

        for col_idx, value in enumerate(member, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = font_leader if is_leader and col_idx in (2, 5) else font_data
            if use_zebra:
                cell.fill = fill_zebra
            cell.border = thin_border
            
            if col_idx in (1, 3, 4):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Footnote
    foot_row = start_row + len(members) + 2
    ws.merge_cells(f"A{foot_row}:F{foot_row}")
    ws[f"A{foot_row}"] = "* Ghi chú: Mọi thành viên trong nhóm đều tham gia đóng góp 100% trách nhiệm và đồng thuận với kết quả nghiệm thu đồ án."
    ws[f"A{foot_row}"].font = Font(name="Calibri", size=10, italic=True, color="666666")
    ws[f"A{foot_row}"].alignment = Alignment(horizontal="left", vertical="center")

    # Column widths
    col_widths = {1: 8, 2: 24, 3: 15, 4: 10, 5: 32, 6: 55}
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[OK] Da tao file Excel tai: {output_path}")

if __name__ == "__main__":
    dest = Path("draft/fraud-detection/reports/Danh sách nhóm.xlsx")
    create_team_excel(dest)
